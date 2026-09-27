import io

from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.routers.formatar import DOCX_MIME
from app.services.normalizacao_fametro import normalizar_documento


def documento_exemplo() -> Document:
    documento = Document()
    for texto in (
        "Título do artigo",
        "Nome do autor",
        "RESUMO",
        "Síntese do trabalho.",
        "Palavras-chave: teste; artigo; sistema.",
        "1. INTRODUÇÃO",
        "Texto da introdução.",
        "2. REFERENCIAL TEÓRICO",
        "2.1 Conceito principal",
        "Texto da fundamentação.",
        "REFERÊNCIAS",
        "AUTOR. Título da obra. Manaus, 2026.",
    ):
        documento.add_paragraph(texto)
    return documento


def bytes_do_documento(documento: Document) -> bytes:
    arquivo = io.BytesIO()
    documento.save(arquivo)
    return arquivo.getvalue()


def test_servico_classifica_e_normaliza_o_artigo():
    documento = documento_exemplo()
    textos_antes = [paragrafo.text for paragrafo in documento.paragraphs]

    totais = normalizar_documento(documento)

    assert [paragrafo.text for paragrafo in documento.paragraphs] == textos_antes
    assert documento.paragraphs[0].style.name == "FAMETRO Título"
    assert documento.paragraphs[1].style.name == "FAMETRO Autor"
    assert documento.paragraphs[3].style.name == "FAMETRO Resumo"
    assert documento.paragraphs[5].style.name == "FAMETRO Seção"
    assert documento.paragraphs[8].style.name == "FAMETRO Subseção"
    assert documento.paragraphs[11].style.name == "FAMETRO Referência"
    assert totais["Seção"] == 2
    assert totais["Subseção"] == 1

    secao = documento.sections[0]
    assert round(secao.top_margin.cm, 1) == 3.0
    assert round(secao.right_margin.cm, 1) == 2.0
    assert round(secao.bottom_margin.cm, 1) == 2.0
    assert round(secao.left_margin.cm, 1) == 3.0


def test_separa_titulos_numerados_colados_ao_corpo():
    documento = Document()
    documento.add_paragraph("Título do artigo")
    documento.add_paragraph("Autor")
    documento.add_paragraph("RESUMO")
    documento.add_paragraph("Resumo do trabalho.")
    documento.add_paragraph("Palavras-chave: teste.")
    documento.add_paragraph(
        "1. INTRODUÇÃO   Texto da introdução."
        "2. REFERENCIAL TEÓRICO   "
    )
    documento.add_paragraph(
        "2.1 Sistemas embarcados   Texto da subseção com "
        "ênfase em um trecho."
    ).runs[-1].italic = True
    documento.add_paragraph("REFERÊNCIAS")
    documento.add_paragraph("AUTOR. Título. Manaus, 2026.")

    normalizar_documento(documento)

    assert [p.text for p in documento.paragraphs[5:10]] == [
        "1. INTRODUÇÃO",
        "Texto da introdução.",
        "2. REFERENCIAL TEÓRICO",
        "2.1 Sistemas embarcados",
        "Texto da subseção com ênfase em um trecho.",
    ]
    assert [p.style.name for p in documento.paragraphs[5:10]] == [
        "FAMETRO Seção",
        "FAMETRO Corpo",
        "FAMETRO Seção",
        "FAMETRO Subseção",
        "FAMETRO Corpo",
    ]
    assert documento.paragraphs[9].runs[0].italic is True
    assert documento.paragraphs[10].style.name == "FAMETRO Título de referências"
    assert documento.paragraphs[11].style.name == "FAMETRO Referência"


def test_endpoint_devolve_docx_normalizado():
    cliente = TestClient(app)
    conteudo = bytes_do_documento(documento_exemplo())

    resposta = cliente.post(
        "/formatar/fametro",
        files={"file": ("artigo científico.docx", conteudo, DOCX_MIME)},
    )

    assert resposta.status_code == 200
    assert resposta.headers["content-type"].startswith(DOCX_MIME)
    assert "artigo_cientfico_fametro.docx" in resposta.headers["content-disposition"]
    assert "filename*=UTF-8''artigo%20cient%C3%ADfico_fametro.docx" in resposta.headers[
        "content-disposition"
    ]

    resultado = Document(io.BytesIO(resposta.content))
    assert resultado.paragraphs[0].style.name == "FAMETRO Título"
    assert resultado.paragraphs[8].style.name == "FAMETRO Subseção"
    assert resultado.paragraphs[11].style.name == "FAMETRO Referência"


def test_endpoint_recusa_arquivo_que_nao_e_docx():
    cliente = TestClient(app)

    resposta = cliente.post(
        "/formatar/fametro",
        files={"file": ("artigo.txt", b"conteudo", "text/plain")},
    )

    assert resposta.status_code == 400
    assert resposta.json()["detail"] == "Apenas arquivos .docx são permitidos."


def test_openapi_documenta_upload_e_download_docx():
    operacao = app.openapi()["paths"]["/formatar/fametro"]["post"]

    assert "multipart/form-data" in operacao["requestBody"]["content"]
    assert DOCX_MIME in operacao["responses"]["200"]["content"]
