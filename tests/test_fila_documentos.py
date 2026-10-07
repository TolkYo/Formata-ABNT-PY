import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fastapi.testclient import TestClient

from app.main import app
from app.routers.formatar import DOCX_MIME
from app.routers import jobs as jobs_router
from app.services import fila_documentos


def _docx_bytes() -> bytes:
    documento = Document()
    documento.add_paragraph("Titulo do artigo")
    documento.add_paragraph("Nome do autor")
    documento.add_paragraph("RESUMO")
    documento.add_paragraph("Sintese do trabalho.")
    documento.add_paragraph("1. INTRODUCAO")
    documento.add_paragraph("Texto da introducao.")
    documento.add_paragraph("REFERENCIAS")
    documento.add_paragraph("AUTOR. Titulo. Manaus, 2026.")
    arquivo = io.BytesIO()
    documento.save(arquivo)
    return arquivo.getvalue()


def test_endpoint_enfileira_e_retorna_202(monkeypatch):
    monkeypatch.setattr(
        fila_documentos,
        "enfileirar",
        lambda **kwargs: {
            "id": "a" * 32,
            "status": "na_fila",
            "posicao": 1,
            "nome_saida": "artigo_formatado.docx",
        },
    )
    cliente = TestClient(app)
    resposta = cliente.post(
        "/jobs/formatar",
        data={"modo": "abnt"},
        files={"file": ("artigo.docx", _docx_bytes(), DOCX_MIME)},
    )

    assert resposta.status_code == 202
    assert resposta.json()["status"] == "na_fila"
    assert resposta.json()["status_url"] == f"/jobs/{'a' * 32}"


def test_endpoint_informa_status_e_download(monkeypatch, tmp_path):
    saida = tmp_path / "saida.docx"
    saida.write_bytes(_docx_bytes())
    removidos = []
    monkeypatch.setattr(
        fila_documentos,
        "status_job",
        lambda job_id: {
            "id": job_id,
            "status": "concluido",
            "posicao": None,
            "mensagem": "Documento pronto para download.",
            "nome_saida": "resultado.docx",
            "download_url": f"/jobs/{job_id}/download",
        },
    )
    monkeypatch.setattr(
        fila_documentos,
        "caminho_download",
        lambda job_id: (saida, "resultado.docx"),
    )
    monkeypatch.setattr(
        fila_documentos,
        "remover_job",
        lambda job_id: removidos.append(job_id),
    )
    cliente = TestClient(app)
    job_id = "b" * 32

    status = cliente.get(f"/jobs/{job_id}")
    download = cliente.get(f"/jobs/{job_id}/download")

    assert status.status_code == 200
    assert status.json()["status"] == "concluido"
    assert download.status_code == 200
    assert download.content == saida.read_bytes()
    assert removidos == [job_id]


def test_worker_processa_em_volume_temporario(monkeypatch, tmp_path):
    monkeypatch.setattr(fila_documentos.settings, "jobs_dir", str(tmp_path))
    monkeypatch.setattr(fila_documentos.metricas, "registrar_evento", lambda *args: None)
    job_id = "c" * 32
    pasta = tmp_path / job_id
    pasta.mkdir()
    (pasta / "entrada.docx").write_bytes(_docx_bytes())

    retorno = fila_documentos.processar_documento(job_id, "abnt", {}, "hash")

    assert retorno == {"arquivo": "saida.docx"}
    assert not (pasta / "entrada.docx").exists()
    resultado = Document(str(pasta / "saida.docx"))
    assert resultado.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert resultado.paragraphs[0].runs[0].font.name == "Arial"


def test_validacao_rejeita_conteudo_falso():
    try:
        fila_documentos.validar_docx_bytes("arquivo.docx", b"nao e zip")
    except ValueError as erro:
        assert "corrompido" in str(erro)
    else:  # pragma: no cover
        raise AssertionError("Conteudo invalido deveria ser recusado")
