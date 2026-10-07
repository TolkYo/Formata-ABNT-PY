"""Testes do formatador ABNT (app/formatter.py)."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from app.formatter import formatar_abnt, formatar_titulos


def _documento():
    doc = Document()
    doc.add_heading("Capitulo 1", level=1)
    doc.add_paragraph("Corpo do texto.")
    doc.add_heading("Subcapitulo", level=2)
    doc.add_heading("Detalhe", level=3)
    return doc


def test_formatar_titulos_classifica_por_estilo():
    doc = _documento()
    formatar_titulos(doc)

    h1 = doc.paragraphs[0]
    assert h1.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert h1.runs[0].bold is True
    assert h1.runs[0].font.name == "Arial"
    assert h1.runs[0].font.size == Pt(14)

    h2 = doc.paragraphs[2]
    assert h2.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert h2.runs[0].font.size == Pt(12)
    assert h2.runs[0].bold is True

    h3 = doc.paragraphs[3]
    assert h3.runs[0].font.size == Pt(12)


def test_corpo_nao_e_tratado_como_titulo():
    doc = _documento()
    formatar_titulos(doc)

    corpo = doc.paragraphs[1]
    assert corpo.runs[0].bold in (None, False)
    assert corpo.runs[0].font.size is None
    assert corpo.alignment != WD_ALIGN_PARAGRAPH.CENTER


def test_formatar_abnt_aplica_margens_e_paragrafos():
    doc = _documento()
    formatar_abnt(doc)

    secao = doc.sections[0]
    # Arredondamento interno (twips) pode variar poucos EMU.
    assert abs(secao.top_margin - Cm(3)) < 5000
    assert abs(secao.left_margin - Cm(3)) < 5000
    assert abs(secao.bottom_margin - Cm(2)) < 5000
    assert abs(secao.right_margin - Cm(2)) < 5000
    for paragrafo in doc.paragraphs:
        assert paragrafo.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
        assert paragrafo.paragraph_format.line_spacing == 1.5
        for run in paragrafo.runs:
            assert run.font.name == "Arial"
            assert run.font.size == Pt(12)
