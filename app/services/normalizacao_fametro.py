"""Normalização de artigos DOCX conforme o modelo FAMETRO."""

from __future__ import annotations

from copy import deepcopy
import re
import unicodedata

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Cm, Pt


PREFIXO = "FAMETRO"
SECAO_RE = re.compile(r"^\s*\d+\s*[.)]\s*(?!\d)")
SUBSECAO_RE = re.compile(r"^\s*\d+\s*[.)]\s*\d+")
SECAO_EMBUTIDA_RE = re.compile(
    r"(?P<titulo>(?<!\d)\d+\.\s+"
    r"[A-ZÀ-ÖØ-Þ][A-ZÀ-ÖØ-Þ0-9 /(),:&–—-]*?)"
    r"(?:\s{2,}|$)"
)
SUBSECAO_COM_CORPO_RE = re.compile(
    r"^(?P<titulo>\d+\.\d+\.?\s+.+?)(?:\s{2,})(?=\S)"
)

# Elementos cujo parágrafo não pode ser repartido com segurança.
_ELEMENTOS_NAO_REPARTIVEIS = frozenset(
    qn(tag)
    for tag in (
        "w:drawing",
        "w:object",
        "w:fldChar",
        "w:footnoteReference",
        "w:hyperlink",
    )
)


def _texto_comparavel(texto: str) -> str:
    sem_acentos = "".join(
        caractere
        for caractere in unicodedata.normalize("NFD", texto)
        if unicodedata.category(caractere) != "Mn"
    )
    return " ".join(sem_acentos.upper().split())


def _estilo_por_id(documento: Document, style_id: str):
    return next(estilo for estilo in documento.styles if estilo.style_id == style_id)


def _obter_ou_criar_estilo(documento: Document, nome: str):
    for estilo in documento.styles:
        if estilo.name == nome:
            return estilo
    return documento.styles.add_style(nome, WD_STYLE_TYPE.PARAGRAPH)


def _configurar_fonte(estilo, tamanho: int, *, negrito: bool = False) -> None:
    estilo.font.name = "Arial"
    estilo.font.size = Pt(tamanho)
    estilo.font.bold = negrito
    estilo._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")


def _configurar_estilos(documento: Document) -> dict[str, object]:
    normal = _estilo_por_id(documento, "Normal")
    _configurar_fonte(normal, 12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.5)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    # tamanho, negrito, alinhamento, recuo, entrelinhas, antes, depois, manter próximo
    especificacoes = {
        "Título": (12, True, WD_ALIGN_PARAGRAPH.CENTER, 0, 1.0, 0, 0, True),
        "Autor": (12, False, WD_ALIGN_PARAGRAPH.RIGHT, 0, 1.0, 0, 0, True),
        "Título do resumo": (12, True, WD_ALIGN_PARAGRAPH.CENTER, 0, 1.0, 0, 0, True),
        "Resumo": (10, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 1.0, 0, 0, False),
        "Palavras-chave": (10, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 1.0, 0, 0, False),
        "Seção": (12, True, WD_ALIGN_PARAGRAPH.LEFT, 0, 1.5, 12, 6, True),
        "Subseção": (12, True, WD_ALIGN_PARAGRAPH.LEFT, 0, 1.5, 6, 0, True),
        "Corpo": (12, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.5, 1.5, 0, 0, False),
        "Título de referências": (12, True, WD_ALIGN_PARAGRAPH.CENTER, 0, 1.5, 12, 6, True),
        "Referência": (12, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 1.0, 0, 12, False),
        "Espaço": (12, False, WD_ALIGN_PARAGRAPH.LEFT, 0, 1.0, 0, 0, False),
    }

    estilos = {}
    existentes = {estilo.name: estilo for estilo in documento.styles}

    def obter(nome: str):
        estilo = existentes.get(nome)
        if estilo is None:
            estilo = documento.styles.add_style(nome, WD_STYLE_TYPE.PARAGRAPH)
            existentes[nome] = estilo
        return estilo

    for rotulo, valores in especificacoes.items():
        tamanho, negrito, alinhamento, recuo, entrelinhas, antes, depois, manter = valores
        estilo = obter(f"{PREFIXO} {rotulo}")
        estilo.base_style = normal
        _configurar_fonte(estilo, tamanho, negrito=negrito)
        formato = estilo.paragraph_format
        formato.alignment = alinhamento
        formato.first_line_indent = Cm(recuo)
        formato.left_indent = Cm(0)
        formato.right_indent = Cm(0)
        formato.line_spacing = entrelinhas
        formato.space_before = Pt(antes)
        formato.space_after = Pt(depois)
        formato.keep_with_next = manter
        estilos[rotulo] = estilo
    return estilos


def _limpar_formatacao_direta(paragrafo) -> None:
    """Deixa o estilo controlar o parágrafo sem apagar ênfases do conteúdo."""
    propriedades = paragrafo._p.get_or_add_pPr()
    preservadas = {qn("w:pStyle"), qn("w:numPr"), qn("w:sectPr")}
    for elemento in list(propriedades):
        if elemento.tag not in preservadas:
            propriedades.remove(elemento)

    for trecho in paragrafo.runs:
        if trecho._r.find(qn("w:footnoteReference")) is not None:
            continue
        propriedades_trecho = trecho._r.rPr
        if propriedades_trecho is None:
            continue
        for nome in ("rFonts", "sz", "szCs"):
            elemento = propriedades_trecho.find(qn(f"w:{nome}"))
            if elemento is not None:
                propriedades_trecho.remove(elemento)
        if trecho.bold is False:
            trecho.bold = None
        if trecho.italic is False:
            trecho.italic = None
        if trecho.underline is False:
            trecho.underline = None


def _intervalo_sem_espacos(texto: str, inicio: int, fim: int) -> tuple[int, int]:
    while inicio < fim and texto[inicio].isspace():
        inicio += 1
    while fim > inicio and texto[fim - 1].isspace():
        fim -= 1
    return inicio, fim


def _partes_estruturais(texto: str) -> list[tuple[int, int]]:
    """Separa títulos numerados colados ao corpo por espaços consecutivos."""
    secoes = list(SECAO_EMBUTIDA_RE.finditer(texto))
    if secoes:
        partes: list[tuple[int, int]] = []
        cursor = 0
        for secao in secoes:
            if secao.start() > cursor:
                partes.append(_intervalo_sem_espacos(texto, cursor, secao.start()))
            partes.append(secao.span("titulo"))
            cursor = secao.end()
        if cursor < len(texto):
            partes.append(_intervalo_sem_espacos(texto, cursor, len(texto)))
        return [parte for parte in partes if parte[0] < parte[1]]

    subsecao = SUBSECAO_COM_CORPO_RE.match(texto)
    if subsecao:
        corpo = _intervalo_sem_espacos(texto, subsecao.end(), len(texto))
        return [subsecao.span("titulo"), corpo]
    return [(0, len(texto))]


def _copiar_trechos(paragrafo, novo: Paragraph, inicio: int, fim: int) -> None:
    posicao = 0
    for trecho in paragrafo.runs:
        texto = trecho.text
        proxima = posicao + len(texto)
        parte_inicio = max(inicio, posicao)
        parte_fim = min(fim, proxima)
        if parte_inicio < parte_fim:
            copia = novo.add_run(texto[parte_inicio - posicao : parte_fim - posicao])
            if trecho._r.rPr is not None:
                copia._r.insert(0, deepcopy(trecho._r.rPr))
        posicao = proxima


def _separar_blocos_concatenados(documento: Document) -> None:
    """Transforma título e corpo colados em parágrafos independentes."""
    for paragrafo in list(documento.paragraphs):
        texto = paragrafo.text
        partes = _partes_estruturais(texto)
        if len(partes) < 2:
            continue

        # Elementos não textuais não podem ser repartidos com segurança.
        if any(el.tag in _ELEMENTOS_NAO_REPARTIVEIS for el in paragrafo._p.iter()):
            continue

        anterior = paragrafo._p
        for inicio, fim in partes:
            elemento = OxmlElement("w:p")
            if paragrafo._p.pPr is not None:
                elemento.append(deepcopy(paragrafo._p.pPr))
            anterior.addnext(elemento)
            novo = Paragraph(elemento, paragrafo._parent)
            _copiar_trechos(paragrafo, novo, inicio, fim)
            anterior = elemento
        paragrafo._p.getparent().remove(paragrafo._p)


def classificar_paragrafos(documento: Document) -> list[str]:
    """Classifica parágrafos pela estrutura textual do artigo."""
    paragrafos = documento.paragraphs
    textos = [p.text for p in paragrafos]
    aparados = [texto.strip() for texto in textos]
    nao_vazios = [indice for indice, texto in enumerate(aparados) if texto]
    if not nao_vazios:
        return ["Espaço"] * len(paragrafos)

    primeiro = nao_vazios[0]
    normalizados = [_texto_comparavel(texto) for texto in textos]
    indices_resumo = [
        indice
        for indice, texto in enumerate(normalizados)
        if texto in {"RESUMO", "ABSTRACT"}
    ]
    primeiro_resumo = indices_resumo[0] if indices_resumo else None
    classificacoes = ["Espaço" if not texto else "Corpo" for texto in aparados]
    classificacoes[primeiro] = "Título"

    if primeiro_resumo is not None:
        for indice in range(primeiro + 1, primeiro_resumo):
            if aparados[indice]:
                classificacoes[indice] = "Autor"

    em_resumo = False
    em_referencias = False
    for indice in range(len(paragrafos)):
        texto = aparados[indice]
        comparavel = normalizados[indice]
        if not texto or indice == primeiro:
            continue
        if comparavel in {"RESUMO", "ABSTRACT"}:
            classificacoes[indice] = "Título do resumo"
            em_resumo = True
            continue
        if comparavel.startswith(("PALAVRAS-CHAVE:", "PALAVRAS CHAVE:", "KEYWORDS:")):
            classificacoes[indice] = "Palavras-chave"
            em_resumo = False
            continue
        if comparavel in {"REFERENCIAS", "REFERENCES"}:
            classificacoes[indice] = "Título de referências"
            em_resumo = False
            em_referencias = True
            continue
        if em_referencias:
            classificacoes[indice] = "Referência"
        elif em_resumo:
            classificacoes[indice] = "Resumo"
        elif SUBSECAO_RE.match(texto):
            classificacoes[indice] = "Subseção"
        elif SECAO_RE.match(texto):
            classificacoes[indice] = "Seção"

    return classificacoes


def normalizar_documento(documento: Document) -> dict[str, int]:
    """Normaliza o documento em memória e retorna contagens por estilo aplicado."""
    _separar_blocos_concatenados(documento)
    estilos = _configurar_estilos(documento)

    for secao in documento.sections:
        secao.orientation = WD_ORIENT.PORTRAIT
        secao.page_width = Cm(21)
        secao.page_height = Cm(29.7)
        secao.top_margin = Cm(3)
        secao.left_margin = Cm(3)
        secao.bottom_margin = Cm(2)
        secao.right_margin = Cm(2)

    ids_estilo = {rotulo: estilo.style_id for rotulo, estilo in estilos.items()}
    classificacoes = classificar_paragrafos(documento)
    totais = {rotulo: 0 for rotulo in estilos}
    for paragrafo, rotulo in zip(documento.paragraphs, classificacoes):
        _limpar_formatacao_direta(paragrafo)
        paragrafo._p.style = ids_estilo[rotulo]
        totais[rotulo] += 1
    return totais
