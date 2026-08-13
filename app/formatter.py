from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def formatar_abnt(doc: Document):
    """
    Aplica as principais regras ABNT diretamente no objeto Document.
    Atenção: Não altera o conteúdo textual, apenas estilos e layout.
    """

    # 1. Configuração das Margens (Superior/Esquerda: 3cm, Inferior/Direita: 2cm)
    secao = doc.sections[0]
    secao.top_margin = Inches(1.18)    # Aproximadamente 3 cm
    secao.left_margin = Inches(1.18)
    secao.bottom_margin = Inches(0.79) # Aproximadamente 2 cm
    secao.right_margin = Inches(0.79)

    # 2. Formatação dos Parágrafos
    for paragrafo in doc.paragraphs:
        # Alinhamento justificado
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        # Espaçamento entre linhas: 1.5
        paragrafo.paragraph_format.line_spacing = 1.5

        # Fonte Arial, tamanho 12 para todos os trechos (runs)
        for run in paragrafo.runs:
            run.font.name = 'Arial'
            run.font.size = Pt(12)


def _nivel_titulo(paragrafo) -> int:
    """Retorna 1, 2 ou 3 conforme o nível do estilo de título, ou 0 se não for título."""
    nome = (paragrafo.style.name or '').lower()
    sid = (paragrafo.style.style_id or '').lower()
    if 'heading 1' in nome or 'título 1' in nome or 'titulo 1' in nome or sid == 'heading1':
        return 1
    if 'heading 2' in nome or 'título 2' in nome or 'titulo 2' in nome or sid == 'heading2':
        return 2
    if 'heading 3' in nome or 'título 3' in nome or 'titulo 3' in nome or sid == 'heading3':
        return 3
    return 0


def formatar_titulos(doc: Document):
    """
    Ajusta os estilos de título (Heading 1/2/3) para as regras ABNT:
    negrito, fonte Arial, tamanhos específicos e alinhamento adequado.
    """
    for paragrafo in doc.paragraphs:
        nivel = _nivel_titulo(paragrafo)
        if nivel == 0:
            continue

        if nivel == 1:
            paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragrafo.paragraph_format.space_before = Pt(18)
            paragrafo.paragraph_format.space_after = Pt(12)
            tamanho = Pt(14)
        elif nivel == 2:
            paragrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragrafo.paragraph_format.space_before = Pt(12)
            paragrafo.paragraph_format.space_after = Pt(6)
            tamanho = Pt(12)
        else:
            paragrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragrafo.paragraph_format.space_before = Pt(6)
            paragrafo.paragraph_format.space_after = Pt(6)
            tamanho = Pt(12)

        for run in paragrafo.runs:
            run.bold = True
            run.font.name = 'Arial'
            run.font.size = tamanho
            run.font.color.rgb = RGBColor(0, 0, 0)


def adicionar_capa(doc: Document, dados: dict):
    """
    Insere uma capa formatada no início do documento usando os dados
    fornecidos (instituicao, curso, autor, titulo, subtitulo, cidade, ano).
    """
    primeiro = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()

    def add(texto, *, negrito=False, tamanho=12):
        p = primeiro.insert_paragraph_before(texto)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.5
        for run in p.runs:
            run.bold = negrito
            run.font.name = 'Arial'
            run.font.size = Pt(tamanho)
        return p

    def linha_vazia(qtd=1):
        for _ in range(qtd):
            p = primeiro.insert_paragraph_before('')
            p.paragraph_format.line_spacing = 1.5

    instituicao = (dados.get('instituicao') or '').strip()
    curso = (dados.get('curso') or '').strip()
    autor = (dados.get('autor') or '').strip()
    titulo = (dados.get('titulo') or '').strip()
    subtitulo = (dados.get('subtitulo') or '').strip()
    cidade = (dados.get('cidade') or '').strip()
    ano = (dados.get('ano') or '').strip()

    if instituicao:
        add(instituicao, negrito=True)
    if curso:
        add(curso)
    linha_vazia(2)

    if autor:
        add(autor, negrito=True)
    linha_vazia(6)

    if titulo:
        add(titulo, negrito=True)
    if subtitulo:
        add(subtitulo)
    linha_vazia(6)

    if cidade:
        add(cidade)
    if ano:
        add(ano)

    # Quebra de página no fim da capa
    quebra = primeiro.insert_paragraph_before('')
    quebra.add_run().add_break(WD_BREAK.PAGE)


def _adicionar_campo_toc(paragrafo):
    """Insere o código de campo XML de um sumário automático (TOC) do Word."""
    run = paragrafo.add_run()

    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')

    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = r'TOC \o "1-3" \h \z \u'

    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')

    texto = OxmlElement('w:t')
    texto.text = 'Atualize o campo no Word (F9 ou "Atualizar campo") para gerar o sumário.'

    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')

    r = run._r
    r.append(fld_begin)
    r.append(instr)
    r.append(fld_sep)
    r.append(texto)
    r.append(fld_end)


def adicionar_sumario(doc: Document):
    """Insere um título e um campo de sumário automático (TOC) no início do documento."""
    primeiro = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()

    titulo = primeiro.insert_paragraph_before('SUMÁRIO')
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo.paragraph_format.line_spacing = 1.5
    for run in titulo.runs:
        run.bold = True
        run.font.name = 'Arial'
        run.font.size = Pt(14)

    campo = primeiro.insert_paragraph_before('')
    campo.paragraph_format.line_spacing = 1.5
    _adicionar_campo_toc(campo)

    quebra = primeiro.insert_paragraph_before('')
    quebra.add_run().add_break(WD_BREAK.PAGE)


SECOES_MINIMAS = {
    'resumo': ('resumo',),
    'abstract': ('abstract',),
    'sumário': ('sumário', 'sumario'),
    'introdução': ('introdução', 'introducao'),
    'conclusão': ('conclusão', 'conclusao'),
    'referências': ('referências', 'referencias'),
}


def validar_estrutura(doc: Document) -> list:
    """
    Verifica se o documento possui a estrutura mínima de um trabalho acadêmico.
    Retorna a lista de seções ausentes (vazia se estiver tudo ok).
    """
    texto = "\n".join(p.text.lower() for p in doc.paragraphs)
    faltantes = []
    for secao, termos in SECOES_MINIMAS.items():
        if not any(termo in texto for termo in termos):
            faltantes.append(secao)
    return faltantes
